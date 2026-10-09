import base64
import io
import json
from datetime import datetime

import httpx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from app.config import get_settings


def _fallback(lines, fields):
    return {
        "issuing_org": fields.get("issuing_org"), "document_number": fields.get("doc_number"),
        "issued_date": str(fields.get("issued_date") or ""), "title": fields.get("title"),
        "recipient": None, "signer": None,
        "body": [line["text"] for line in lines if line.get("text")], "tables": [],
        "uncertain_fields": ["Bản nháp dùng PaddleOCR; Văn thư cần kiểm tra nội dung."],
    }


async def extract_document(images, lines, fields):
    settings = get_settings()
    fallback = _fallback(lines, fields)
    if not settings.llm_enabled or not settings.llm_model:
        return fallback
    content = [{"type": "text", "text": "Extract this Vietnamese administrative document as compact JSON only: issuing_org, document_number, issued_date, title, recipient, signer, body (array of paragraphs), tables (array of {headers,rows}), uncertain_fields. Do not invent unreadable or blank text."}]
    for image in images[:settings.llm_max_pages]:
        image = image.convert("RGB")
        image.thumbnail((1600, 1600))
        buffer = io.BytesIO(); image.save(buffer, "JPEG", quality=85)
        content.append({"type": "image_url", "image_url": {"url": "data:image/jpeg;base64," + base64.b64encode(buffer.getvalue()).decode()}})
    payload = {"model": settings.llm_model, "temperature": 0, "max_tokens": 1800, "reasoning_effort": "none", "messages": [{"role": "user", "content": content}]}
    try:
        async with httpx.AsyncClient(timeout=settings.llm_timeout_seconds) as client:
            response = await client.post(settings.llm_base_url.rstrip("/") + "/chat/completions", json=payload)
            response.raise_for_status()
        value = json.loads(response.json()["choices"][0]["message"]["content"].replace("```json", "").replace("```", ""))
        return {**fallback, **value}
    except Exception:
        fallback["uncertain_fields"].append("LLM không phản hồi; cần kiểm tra thủ công.")
        return fallback


def create_editable_docx(data):
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(2)
    document.styles["Normal"].font.name = "Times New Roman"
    document.styles["Normal"].font.size = Pt(12)
    paragraph = document.add_paragraph("BẢN NHÁP OCR + LLM – Vui lòng kiểm tra trước khi sử dụng")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER; paragraph.runs[0].bold = True
    for key in ("issuing_org", "document_number", "issued_date", "title"):
        value = data.get(key)
        if value:
            paragraph = document.add_paragraph(str(value)); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if key == "title": paragraph.runs[0].bold = True
    for label, key in (("Kính gửi", "recipient"), ("Người ký", "signer")):
        if data.get(key): document.add_paragraph(f"{label}: {data[key]}")
    for line in data.get("body", []):
        if line: document.add_paragraph(str(line))
    for source in data.get("tables", []):
        headers = source.get("headers", [])
        if not headers: continue
        table = document.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"
        for index, value in enumerate(headers): table.rows[0].cells[index].text = str(value)
        for row in source.get("rows", []):
            cells = table.add_row().cells
            for index, value in enumerate(row[:len(cells)]): cells[index].text = str(value)
    document.add_heading("Nội dung cần Văn thư kiểm tra", level=2)
    for item in data.get("uncertain_fields", []): document.add_paragraph(str(item), style="List Bullet")
    document.add_paragraph(f"Xuất lúc {datetime.now():%d/%m/%Y %H:%M}")
    output = io.BytesIO(); document.save(output)
    return output.getvalue()
