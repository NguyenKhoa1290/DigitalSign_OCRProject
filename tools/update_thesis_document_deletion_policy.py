"""Add the approved draft deletion design without rewriting unrelated DOCX parts."""
import argparse
import copy
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree

from update_thesis_document_workflows import NS, W, new_paragraph, replace, text


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--overwrite', action='store_true')
    args = parser.parse_args()
    if args.source.resolve() == args.output.resolve():
        raise ValueError('Preserve the reviewed source')
    if args.output.exists() and not args.overwrite:
        raise ValueError('Output already exists')
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(args.source) as zin:
        root = etree.fromstring(zin.read('word/document.xml'))
        body = root.find(W + 'body')
        original_children = list(body)
        original_bytes = {id(n): etree.tostring(n) for n in original_children}
        paragraphs = [p for p in body.findall(W + 'p')
                      if not str(p.xpath('string(w:pPr/w:pStyle/@w:val)', namespaces=NS)).lower().startswith('toc')]
        tables = body.findall(W + 'tbl')
        normal = next(p for p in paragraphs if text(p).startswith('Phạm vi đề tài giữ nguyên hai luồng'))

        def match(prefix):
            found = [p for p in paragraphs if text(p).startswith(prefix)]
            if len(found) != 1:
                raise ValueError(f'Expected a unique body paragraph: {prefix}')
            return found[0]

        def insert_before(anchor, nodes):
            for node in nodes:
                body.insert(body.index(anchor), node)

        def paras(values):
            return [new_paragraph(normal, value) for value in values]

        def set_cell(cell, values):
            templates = cell.findall(W + 'p')
            template = next((p for p in templates if text(p)), templates[0])
            for p in templates:
                cell.remove(p)
            for value in values:
                cell.append(new_paragraph(template, value))

        # The existing role matrix keeps all its original rows, order and styling.
        matrix = tables[7]
        source_matrix_bytes = [etree.tostring(r) for r in matrix.findall(W + 'tr')]
        template_row = next(r for r in matrix.findall(W + 'tr') if text(r).startswith('Soạn thảo, tải công văn lên'))
        matrix_rows = [
            ['Xóa mềm dự thảo cá nhân', '', 'X (1)', 'X (1)', '', ''],
            ['Xóa mềm dự thảo hỗ trợ', 'X (2)', '', '', '', ''],
            ['Khôi phục dự thảo', 'X (3)', '', '', '', ''],
        ]
        for offset, values in enumerate(matrix_rows, start=1):
            row = copy.deepcopy(template_row)
            for height in row.findall('./w:trPr/w:trHeight', NS):
                height.getparent().remove(height)
            for cell, value in zip(row.findall(W + 'tc'), values):
                set_cell(cell, [value])
            matrix.insert(matrix.index(template_row) + offset, row)
        insert_before(match('2.4. Thiết kế hệ thống và sơ đồ thiết kế hệ thống'), paras([
            'Bổ sung phân quyền ngày 19/10/2026: các quyền (1), (2), (3) là yêu cầu thiết kế cần triển khai, không xác nhận API hiện tại đã đáp ứng. Các quyền khác trong bảng được giữ nguyên.',
            '(1) Văn thư và Chuyên viên chỉ được xóa mềm dự thảo do mình tạo, còn Draft, chưa trình duyệt, chưa phân công xử lý, chưa ký và chưa tiếp nhận chính thức vào sổ công văn đến. Không xác định điều kiện này chỉ bằng trạng thái hiện tại.',
            '(2) Admin được xóa mềm dự thảo của người khác để xử lý nhập sai hoặc trùng, với cùng điều kiện nghiệp vụ và bắt buộc nhập lý do. Vai trò quản trị không cho phép bỏ qua giai đoạn xử lý hoặc xóa công văn đã phát hành.',
            '(3) Admin được khôi phục dự thảo đã xóa mềm khi điều kiện nghiệp vụ còn hợp lệ; phải ghi lý do, người thực hiện và thời điểm. Lãnh đạo Phòng, Ban Giám hiệu không có quyền xóa hoặc khôi phục chỉ từ vai trò lãnh đạo.',
            'Không cho xóa công văn đã trình duyệt, bị trả lại (Rejected), đang ký, đã ký, đã phát hành hoặc công văn đến đã tiếp nhận chính thức. Trường hợp dừng xử lý hoặc thu hồi cần quy trình riêng, giữ lý do và lịch sử; bổ sung này không tự cấp quyền hủy hoặc thu hồi.',
        ]))

        # Add a use case with the same two-column form as the existing scenarios.
        scenario = copy.deepcopy(tables[2])
        rows = scenario.findall(W + 'tr')
        for index, row in enumerate(rows):
            if index == 5:  # The main scenario may span pages; short rows stay intact.
                continue
            props = row.find(W + 'trPr')
            if props is None:
                props = etree.SubElement(row, W + 'trPr')
            if props.find(W + 'cantSplit') is None:
                etree.SubElement(props, W + 'cantSplit')
            if index == 0 and props.find(W + 'tblHeader') is None:
                etree.SubElement(props, W + 'tblHeader')
        scenario_values = [
            ['Xóa mềm và khôi phục bản nháp (SoftDeleteAndRestoreDraft)'],
            ['Chuyên viên, Văn thư; Admin thực hiện hỗ trợ và khôi phục.'],
            ['Người dùng đã đăng nhập và có quyền. Hồ sơ còn Draft, chưa trình duyệt, phân công, ký hoặc tiếp nhận chính thức. Chuyên viên/Văn thư là người tạo; Admin hỗ trợ bắt buộc có lý do.'],
            ['Xóa: hồ sơ được đánh dấu xóa mềm và ẩn khỏi danh sách thông thường; tệp, chữ ký và lịch sử không bị xóa vật lý. Khôi phục: hồ sơ hợp lệ trở lại danh sách, giữ nguyên lịch sử và có nhật ký khôi phục.'],
            ['1. Người dùng mở dự thảo hợp lệ và chọn Xóa bản nháp; Admin có thể chọn hồ sơ cần hỗ trợ.',
             '2. Hệ thống hiển thị xác nhận, giải thích xóa mềm; Admin phải nhập lý do.',
             '3. API lấy danh tính từ JWT, kiểm tra quyền, người tạo, lịch sử, tình trạng tiếp nhận và ký.',
             '4. Trong một giao dịch, API kiểm tra phiên bản, đánh dấu xóa mềm và ghi người xóa, thời điểm, lý do cùng nhật ký.',
             '5. Hệ thống thông báo kết quả và cập nhật danh sách; không xóa tệp MinIO hoặc lịch sử luân chuyển.',
             '6. Admin nhập lý do khôi phục; API kiểm tra lại quyền, dữ liệu và ghi nhật ký.'],
            ['Không có quyền hoặc không phải người tạo: từ chối, không thay đổi hồ sơ.',
             'Hồ sơ đã trình duyệt, Rejected, đã phân công, ký, tiếp nhận chính thức hoặc phát hành: không cho xóa, kể cả Admin.',
             'Hồ sơ thay đổi đồng thời, thiếu lý do bắt buộc hoặc không đủ bằng chứng về người tạo/giai đoạn: từ chối và yêu cầu kiểm tra lại.',
             'Giao dịch lỗi: giữ nguyên hồ sơ; tác vụ OCR và callback không được cập nhật hoặc làm sống lại hồ sơ đã xóa.'],
        ]
        for row, values in zip(rows[1:], scenario_values):
            set_cell(row.findall(W + 'tc')[1], values)
        for p in scenario.findall('.//w:p', NS):
            props = p.find(W + 'pPr')
            if props is None:
                props = etree.Element(W + 'pPr')
                p.insert(0, props)
            keep = props.find(W + 'keepLines')
            if keep is None:
                keep = etree.SubElement(props, W + 'keepLines')
            keep.set(W + 'val', '1')
        scenario_heading = new_paragraph(match('2.1.4.6. Kịch bản'), '2.1.4.7. Kịch bản xóa mềm và khôi phục bản nháp')
        insert_before(match('2.2. Mô hình Phân cấp'), [scenario_heading, scenario])

        insert_before(match('2.4.4.3. OCR Service'), paras([
            'Thiết kế bổ sung về xóa mềm: Gateway xác thực, Document Service quyết định quyền theo vai trò, người tạo và toàn bộ giai đoạn xử lý. Nút trên giao diện chỉ hỗ trợ thao tác; API phải kiểm tra lại và không tin người tạo hoặc vai trò do client gửi. Chuyên viên/Văn thư chỉ xóa dự thảo cá nhân đủ điều kiện; Admin hỗ trợ và khôi phục có lý do, không vượt điều kiện nghiệp vụ.',
            'Xóa mềm được thực hiện trong giao dịch có kiểm soát đồng thời với thao tác trình duyệt, phân công, tiếp nhận và ký. Lưu dấu xóa riêng, giữ trạng thái Draft và toàn bộ tệp, lịch sử, bằng chứng ký. Việc hủy xử lý hoặc thu hồi công văn sau trình duyệt thuộc quy trình riêng, không dùng thao tác xóa.',
            'Các API đọc, tải tệp, sửa, OCR, phân công, trình duyệt và ký phải kiểm tra dấu xóa; danh sách, tìm kiếm, thống kê và thông báo thông thường không trả hồ sơ đã xóa. Callback hoặc tác vụ nền đang chạy không được phục hồi hay ghi đè hồ sơ này. Chỉ Admin được truy cập luồng quản lý dự thảo đã xóa để kiểm tra và khôi phục có nhật ký.',
        ]))

        insert_before(match('2.6.3.6. Bảng DocumentTypes'), paras([
            'Thiết kế dữ liệu bổ sung phục vụ xóa mềm (chưa có migration áp dụng): đề xuất lưu CreatedById và CreatedAt bất biến; FirstSubmittedAt đánh dấu lần trình duyệt đầu tiên; OfficiallyReceivedAt đánh dấu công văn đến đã tiếp nhận chính thức. Các mốc này cùng lịch sử phân công và bằng chứng ký quyết định điều kiện xóa, không suy ra chỉ từ giá trị Draft.',
            'Đề xuất bổ sung IsDeleted, DeletedAt, DeletedById và DeleteReason; nhật ký ghi riêng sự kiện SoftDelete và Restore, người thực hiện, thời điểm và lý do. Dùng ConcurrencyVersion hoặc cơ chế kiểm soát đồng thời tương đương. Khôi phục chỉ bỏ dấu xóa sau kiểm tra, không xóa sự kiện lịch sử. Các trường nêu trên là thiết kế bổ sung, chưa phải cấu trúc database đã triển khai.',
            'Giữ liên kết Documents với DocumentProcess, chữ ký và tệp MinIO khi xóa mềm; không gọi xóa bản ghi hoặc cascade delete. Dữ liệu cũ thiếu người tạo/mốc nghiệp vụ cần đối soát trước khi cấp quyền xóa. Code hiện dùng cùng action Submit cho khởi tạo và trình duyệt, vì vậy migration phải phân biệt hai sự kiện; trường hợp không xác định chắc chắn phải chặn xóa.'],
        ))

        end = body.find(W + 'sectPr')
        insert_before(end, [new_paragraph(match('2.8.1. Phạm vi hai luồng'),
            '2.8.4. Phạm vi bổ sung phân quyền xóa công văn', page_break=True)] + paras([
            'Cập nhật thiết kế ngày 19/10/2026 bổ sung quyền xóa mềm và khôi phục bản nháp tại mục 2.1.4.7, ma trận 2.3, Document Service và dữ liệu Documents. Chính sách chỉ cho xử lý dự thảo chưa đi vào quy trình, phân biệt quyền cá nhân của Văn thư/Chuyên viên với hỗ trợ có lý do của Admin. Hồ sơ đã trình, bị trả lại, phân công, ký, tiếp nhận chính thức hoặc phát hành không được xóa bằng chức năng này.',
            'Đối chiếu mã nguồn hiện tại: API còn xóa cứng ở Draft hoặc Rejected, chưa kiểm tra người tạo/vai trò tại thao tác xóa; giao diện chưa có nút xóa. Vì vậy, chính sách mới là thiết kế cần triển khai, không phải kết quả nghiệm thu. Cần bổ sung dữ liệu người tạo, mốc nghiệp vụ và xóa mềm; kiểm soát đồng thời; bổ sung khôi phục và kiểm tra dấu xóa trong truy vấn, thống kê, tải tệp, ký và callback OCR.',
            'Nghiệm thu phải kiểm tra vai trò và người tạo, từ chối mọi giai đoạn không hợp lệ, bảo toàn tệp/lịch sử, khôi phục có lý do và xung đột xóa với trình duyệt hoặc ký. Dữ liệu cũ dùng chung action Submit cho khởi tạo và trình duyệt phải được đối soát trước; không đủ bằng chứng thì chặn xóa. Bộ tiêu chí TC-DOC-DELETE-POLICY-087 mới được lập, chưa chạy vì chức năng chưa được triển khai. Quy trình hủy, thu hồi và thời hạn xóa vật lý cần quyết định riêng, không được coi là đã bổ sung trong hạng mục này.',
        ]))

        # Everything outside the specifically extended matrix stays byte-identical.
        for child in original_children:
            if child is not matrix:
                assert etree.tostring(child) == original_bytes[id(child)], 'Unrelated source content changed'
        remaining_rows = [etree.tostring(r) for r in matrix.findall(W + 'tr') if not text(r).startswith(('Xóa mềm dự thảo', 'Khôi phục dự thảo'))]
        assert remaining_rows == source_matrix_bytes
        xml = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
        with ZipFile(args.output, 'w') as zout:
            for entry in zin.infolist():
                zout.writestr(entry, xml if entry.filename == 'word/document.xml' else zin.read(entry.filename))
    with ZipFile(args.source) as zin, ZipFile(args.output) as zout:
        assert all(zin.read(n) == zout.read(n) for n in zin.namelist() if n != 'word/document.xml')
    doc = Document(args.output)
    assert len(doc.tables) == 17
    assert len(doc.tables[8].rows) == 17  # newly inserted use case precedes the role matrix
    resource = args.output.parents[2] / 'tai_nguyen' / 'noi_dung_trich_xuat'
    resource.mkdir(parents=True, exist_ok=True)
    (resource / 'doc_content.txt').write_text('\n'.join(p.text for p in doc.paragraphs), encoding='utf-8')
    (resource / 'tables_content.txt').write_text('\n\n'.join(
        f'=== TABLE {i} ===\n' + '\n'.join(' | '.join(c.text.strip() for c in row.cells).rstrip() for row in table.rows)
        for i, table in enumerate(doc.tables)) + '\n', encoding='utf-8')
    source_hash = hashlib.sha256(args.source.read_bytes()).hexdigest()
    output_hash = hashlib.sha256(args.output.read_bytes()).hexdigest()
    (resource / 'source.json').write_text(json.dumps({
        'source': args.output.as_posix(), 'source_sha256': output_hash,
        'updated_date': '2026-10-19', 'tables': len(doc.tables), 'paragraphs': len(doc.paragraphs),
    }, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    proof = {
        'source': args.source.as_posix(), 'output': args.output.as_posix(),
        'source_sha256': source_hash, 'output_sha256': output_hash,
        'source_original_children_preserved_except_extended_matrix': True,
        'existing_matrix_rows_preserved': True, 'other_package_parts_preserved': True,
        'tables_before': 16, 'tables_after': len(doc.tables), 'matrix_added_rows': 3,
        'added_scenario': '2.1.4.7', 'implementation_changed': False,
    }
    qa = resource.parent / 'kiem_tra_bao_cao' / 'phan_quyen_xoa_20261019'
    qa.mkdir(parents=True, exist_ok=True)
    (qa / 'verification.json').write_text(json.dumps(proof, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(proof, ensure_ascii=True))


if __name__ == '__main__':
    main()
