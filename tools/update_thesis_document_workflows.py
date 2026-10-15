"""Edit the reviewed thesis while retaining existing OOXML formatting and package parts."""
import argparse
import copy
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
W = "{" + NS["w"] + "}"


def text(p):
    return "".join(p.xpath(".//w:t/text()", namespaces=NS))


def replace(p, value):
    nodes = p.findall(".//w:t", NS)
    if not nodes:
        raise ValueError("Expected a paragraph containing text")
    nodes[0].text = value
    nodes[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    for node in nodes[1:]:
        node.text = ""


def new_paragraph(template, value, page_break=False):
    p = copy.deepcopy(template)
    p.attrib.clear()
    ppr = p.find(W + "pPr")
    first_run = p.find(W + "r")
    rpr = copy.deepcopy(first_run.find(W + "rPr")) if first_run is not None else None
    for child in list(p):
        if child is not ppr:
            p.remove(child)
    if ppr is not None:
        for tag in ("pageBreakBefore", "sectPr"):
            for node in ppr.findall(W + tag):
                ppr.remove(node)
        if page_break:
            etree.SubElement(ppr, W + "pageBreakBefore")
    run = etree.SubElement(p, W + "r")
    if rpr is not None:
        run.append(rpr)
    etree.SubElement(run, W + "t").text = value
    return p


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--overwrite", action="store_true", help="Replace only the task's unreviewed output during layout QA")
    args = parser.parse_args()
    if args.output.resolve() == args.source.resolve():
        raise ValueError("Do not overwrite the source report")
    if args.output.exists() and not args.overwrite:
        raise ValueError("Output already exists; preserve reviewed versions")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(args.source) as source:
        root = etree.fromstring(source.read("word/document.xml"))
        before = copy.deepcopy(root)
        body = root.find(W + "body")
        paragraphs = body.findall(W + "p")
        tables = body.findall(W + "tbl")
        changed = []

        def change(p, value):
            changed.append(root.getroottree().getpath(p))
            replace(p, value)

        def body_match(prefix):
            matches = [p for p in paragraphs if text(p).startswith(prefix)]
            if len(matches) != 1:
                raise ValueError(f"Expected unique body text: {prefix}")
            return matches[0]

        change(body_match("Bảng này lưu trữ danh mục các loại văn bản hành chính"),
               "Bảng này lưu trữ danh mục phân loại văn bản phục vụ tạo lập, tra cứu và quản lý hồ sơ. Phiên bản thử nghiệm sử dụng bốn giá trị: CongVanDen (công văn đến), CongVanDi (công văn đi), ToTrinh (Tờ trình) và QuyetDinh (Quyết định). Danh mục này chưa tự quyết định quyền truy cập hoặc chuyển trạng thái nghiệp vụ.")
        change(body_match("Dữ liệu mẫu (Seed Data): Khi triển khai, bảng này"),
               "Dữ liệu mẫu (Seed Data): Bốn giá trị được thiết lập sẵn để người dùng chọn loại văn bản khi tạo mới, sửa thông tin hoặc lọc danh sách. Công văn đến/đi biểu thị hướng luân chuyển; Tờ trình và Quyết định biểu thị thể loại. Tờ trình hoặc Quyết định được xử lý theo nguồn tiếp nhận và mục đích ban hành, không mặc định hình thành hai luồng nghiệp vụ riêng. Việc tách hướng luân chuyển thành trường riêng là hướng hoàn thiện thiết kế; cơ sở dữ liệu hiện tại vẫn sử dụng một trường DocTypeId.")

        table = tables[4]
        rows = table.findall(W + "tr")
        def cell_paragraphs(row):
            return rows[row].findall(W + "tc")[1].findall(W + "p")
        if "FinalApprovalAndLegalSignature" not in text(rows[1]):
            raise ValueError("Unexpected approval table")
        change(cell_paragraphs(4)[0],
               "Đối với công văn đến: ý kiến chỉ đạo được lưu và văn bản được chuyển tới đơn vị xử lý. Đối với công văn đi: PDF có chữ ký pháp nhân được lưu vào MinIO, chuyển sang trạng thái chờ Văn thư phát hành. Văn bản chỉ được ghi nhận Đã ban hành sau thao tác phát hành.")
        scenario = [
            "1. Ban Giám hiệu truy cập vào không gian làm việc dành cho lãnh đạo.",
            "2. Hệ thống hiển thị văn bản đến cần chỉ đạo và văn bản đi chờ phê duyệt.",
            "3. Chọn văn bản và nhánh đến hoặc đi.",
            "4. Hiển thị PDF và lịch sử ký/duyệt.",
            "5. Với công văn đến: nhập bút phê, chọn đơn vị nhận và nhấn Lưu chỉ đạo; hệ thống lưu ý kiến, chuyển đơn vị xử lý và kết thúc nhánh này.",
            "6. Với công văn đi: kiểm tra dự thảo đã được duyệt cấp phòng và chọn Ký số pháp nhân; các bước 7 đến 10 chỉ áp dụng cho nhánh công văn đi.",
            "7. Hệ thống kiểm tra quyền ký, chứng thư và yêu cầu xác thực theo cơ chế ký được triển khai.",
            "8. Sign Service ký PDF, lưu bản đã ký và trả kết quả; bản gốc và bản đã ký phải được quản lý riêng.",
            "9. Chỉ khi ký thành công, hệ thống ghi nhận Ban Giám hiệu đã ký và chuyển Văn thư phát hành. Lỗi ký không làm thay đổi trạng thái thành đã ký.",
            "10. Văn thư kiểm tra bản đã ký, thực hiện phát hành; hệ thống cập nhật Đã ban hành, ghi lịch sử và lưu trữ."
        ]
        assert len(cell_paragraphs(5)) == len(scenario)
        for p, value in zip(cell_paragraphs(5), scenario):
            change(p, value)
        for p in cell_paragraphs(6):
            value = text(p)
            if value.startswith("6.1. Ban Giám hiệu"):
                change(p, "6.1. Đối với công văn đi, Ban Giám hiệu không đồng ý và yêu cầu chỉnh sửa:")
            elif value.startswith("6.1.1. Ban Giám hiệu"):
                change(p, "6.1.1. Nhập lý do yêu cầu sửa và chọn Trả lại.")
            elif value.startswith("6.2. Lỗi kết nối"):
                change(p, "6.2. Lỗi ký công văn đi hoặc chứng thư số không hợp lệ:")
            elif value.startswith("6.3. Văn bản"):
                change(p, "6.3. Công văn đi chưa có đủ chữ ký kiểm soát của cấp phòng:")

        change(body_match("Phương án đã được lựa chọn và kiểm thử bằng script; chưa tích hợp pipeline"),
               "Phương án đã được lựa chọn, thử nghiệm bằng script và bổ sung endpoint xuất Word bằng PaddleOCR kết hợp Gemini cùng giao diện chọn model. Admin có thể cấu hình API key; kết quả Word là bản nháp để Văn thư kiểm tra. Việc lưu Word gắn với hồ sơ công văn, quản lý phiên bản chỉnh sửa và soạn thảo trong ứng dụng vẫn cần hoàn thiện. Chất lượng bố cục và luồng sử dụng với API key, tài liệu thật cần được nghiệm thu; kết quả trên một mẫu không đại diện cho mọi công văn.")

        section = [
            ("heading2", "2.8. Đối chiếu yêu cầu nghiệp vụ và phạm vi triển khai", True),
            ("heading3", "2.8.1. Phạm vi hai luồng công văn", False),
            ("normal", "Phạm vi đề tài giữ nguyên hai luồng công văn đến và công văn đi. Công văn đến được tiếp nhận, số hóa, đối soát thông tin, chuyển lãnh đạo chỉ đạo, phân công chuyên viên và theo dõi kết quả xử lý. Công văn đi được lập dự thảo, trình duyệt cấp phòng, phê duyệt và ký số, sau đó Văn thư phát hành và lưu trữ. Đây là yêu cầu thiết kế; mức độ đã triển khai được xác định riêng dưới đây.", False),
            ("normal", "Tờ trình và Quyết định được quản lý trong danh mục loại văn bản, dùng các chức năng tạo lập, tra cứu và xử lý hồ sơ chung. Đề tài chưa đặt ra một quy trình độc lập cho từng thể loại. Hướng luân chuyển và nguồn gốc văn bản là căn cứ để xác định nhánh xử lý khi hoàn thiện workflow.", False),
            ("heading3", "2.8.2. Chức năng đã triển khai", False),
            ("normal", "Tại thời điểm rà soát ngày 13/10/2026, hệ thống có danh mục loại văn bản, tạo và sửa metadata, lưu PDF, nhận diện OCR, phân công và lịch sử xử lý. Bộ lọc phân biệt đúng Công văn đến, Công văn đi, Tờ trình và Quyết định; tên loại được hiển thị bằng tiếng Việt.", False),
            ("normal", "Document Service hỗ trợ chuyển trạng thái trình duyệt, duyệt cấp phòng, trình Ban Giám hiệu và phát hành. Sign Service cung cấp ký và xác minh PDF. Hai nhóm thao tác còn tách rời; trạng thái đã ký chưa tự xác nhận PDF được ký thành công. Các loại văn bản hiện dùng chung workflow.", False),
            ("heading3", "2.8.3. Giới hạn và công việc cần hoàn thiện", True),
            ("normal", "Luồng công văn đến cần hoàn thiện thao tác bút phê riêng, chuyển đơn vị, quản lý thời hạn, cập nhật tiến độ và báo cáo kết quả, cùng việc xác nhận kết thúc xử lý. Phân công hiện ghi người giao, người nhận và ý kiến vào lịch sử; chức năng này chưa thay thế toàn bộ quy trình theo dõi nhiệm vụ.", False),
            ("normal", "Luồng công văn đi cần liên kết kết quả ký PDF với trạng thái duyệt, kiểm tra quyền hạn tại backend và bảo đảm đúng phiên bản tài liệu trước khi phát hành. Khi lỗi ký hoặc thay đổi nội dung, hệ thống phải giữ trạng thái phù hợp và có lịch sử để đối soát. Các bước này thuộc yêu cầu cần hoàn thiện, chưa được ghi nhận là nghiệm thu đầy đủ.", False),
            ("normal", "Kiểm thử kỹ thuật bộ lọc ngày 13/10/2026 tại giao diện 1366px và 390px xác nhận tham số lọc, phân trang và bốn lựa chọn loại văn bản hoạt động đúng với dữ liệu mô phỏng. API thật trả dữ liệu đúng loại. Kiểm thử bộ lọc không chứng minh hai luồng nghiệp vụ đã hoàn thành; người dùng cần nghiệm thu trên công văn thật với các vai trò Văn thư, Chuyên viên, Lãnh đạo Phòng và Ban Giám hiệu.", False),
            ("normal", "Kết quả nghiệm thu cần ghi rõ tài liệu đầu vào, người thực hiện, từng bước xử lý, tệp đầu ra, ngày chạy và lỗi còn tồn tại. Đối với OCR, cần đối chiếu Word nháp với bản gốc; đối với ký số, cần xác minh PDF đã ký và trạng thái hồ sơ. Căn cứ kỹ thuật của lần rà soát được lưu tại testcase TC-DOC-TYPE-FILTER-069 và kế hoạch triển khai của dự án.", False),
        ]
        templates = {"heading2": body_match("2.7. Phương án"),
                     "heading3": body_match("2.7.1. Cơ sở"),
                     "normal": paragraphs[525]}
        added = []
        section_properties = body.find(W + "sectPr")
        for style, value, page_break in section:
            p = new_paragraph(templates[style], value, page_break)
            body.insert(body.index(section_properties), p)
            added.append(p)

        # Check the source paragraphs and structural properties outside the edited scope.
        before_paragraphs = before.findall(".//w:p", NS)
        after_existing = [p for p in root.findall(".//w:p", NS) if p not in added]
        assert len(before_paragraphs) == len(after_existing)
        changed_set = set(changed)
        for a, b in zip(before_paragraphs, after_existing):
            if root.getroottree().getpath(b) not in changed_set:
                assert etree.tostring(a) == etree.tostring(b), "Unrelated paragraph changed"
            else:
                assert etree.tostring(a.find(W + "pPr")) == etree.tostring(b.find(W + "pPr")), "Paragraph format changed"
        updated = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        with ZipFile(args.output, "w") as target:
            for entry in source.infolist():
                target.writestr(entry, updated if entry.filename == "word/document.xml" else source.read(entry.filename))
    with ZipFile(args.source) as source, ZipFile(args.output) as target:
        preserved = all(source.read(name) == target.read(name) for name in source.namelist() if name != "word/document.xml")
    assert preserved
    doc = Document(args.output)
    assert len(doc.tables) == 16
    assert [len(table.rows) for table in doc.tables] == [len(table.rows) for table in Document(args.source).tables]
    verification = {"source": str(args.source), "output": str(args.output),
                    "source_sha256": hashlib.sha256(args.source.read_bytes()).hexdigest(),
                    "output_sha256": hashlib.sha256(args.output.read_bytes()).hexdigest(),
                    "modified_paragraphs": len(changed), "added_paragraphs": len(added),
                    "existing_paragraph_format_preserved": True,
                    "unrelated_paragraphs_preserved": True,
                    "other_package_parts_preserved": preserved,
                    "table_count_and_row_counts_preserved": True}
    args.output.with_suffix(".verification.json").write_text(json.dumps(verification, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(verification, ensure_ascii=False))


if __name__ == "__main__":
    main()
