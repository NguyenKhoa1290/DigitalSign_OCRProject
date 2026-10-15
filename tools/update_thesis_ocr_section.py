"""Append the approved OCR design without rewriting existing report contents."""
import argparse
import copy
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree


PAGES = [
    [
        ("Heading2", "2.7. Phương án số hóa tài liệu kết hợp PaddleOCR và Gemini"),
        ("Heading3", "2.7.1. Cơ sở lựa chọn phương án"),
        ("Normal", "Đồ án lựa chọn kết hợp PaddleOCR/PP-StructureV3 với Gemini 3.8 Flash để tái tạo tài liệu thành Word có thể chỉnh sửa. Paddle cung cấp văn bản nhận diện, vị trí vùng chữ và cấu trúc bảng; Gemini đối chiếu ảnh gốc với dữ liệu này để khôi phục nội dung và bố cục. Bộ xuất DOCX local chuyển JSON thành đoạn văn, vùng song song và bảng thực, thay vì đặt ảnh vào trang Word."),
        ("Normal", "Trên mẫu kế hoạch thực tế, Gemini chỉ đọc ảnh đã tạo được Word nhưng có lần đếm 10 hoặc 8 hàng trống thay vì 7. Kết quả PP-StructureV3 đã lưu có 8 hàng tổng, gồm 1 hàng tiêu đề và 7 hàng thân trống. Khi cung cấp cả HTML bảng và tọa độ ô Paddle cùng ảnh, lượt thử với prompt tiếng Việt đã xuất Word đúng 4 cột và 7 hàng trống. Người dùng chấp nhận phương án kết hợp này ngày 12/10/2026."),
        ("Heading3", "2.7.2. Quy trình xử lý được lựa chọn"),
        ("Normal", "Quy trình gồm: tiếp nhận ảnh gốc; nhận diện bằng Paddle/PP-Structure; gửi ảnh cùng chữ, độ tin cậy, tọa độ và HTML bảng cho Gemini; nhận JSON bố cục theo schema; kiểm tra kiểu dữ liệu và số ô/cột; xuất Word local; chuyển bản nháp cho văn thư đối chiếu, sửa lỗi và xác nhận."),
        ("Normal", "Prompt tiếng Việt quy định ảnh là nguồn chính, không tự điền ô trống hoặc tạo chữ ký. Model phải tách hàng tiêu đề khỏi hàng thân, đếm theo đường biên, kiểm tra từ trên xuống và dưới lên, bảo đảm rows cộng blank_rows không bị tính trùng. Không đưa số hàng đúng của một mẫu vào prompt để áp dụng cho mọi công văn."),
        ("Normal", "Paddle hỗ trợ bằng cấu trúc phát hiện thực tế, không phải khuôn cố định. Nội dung chưa chắc chắn được ghi riêng trong uncertain_fields; tên trường JSON giữ nguyên để tương thích bộ xuất Word. Một lần sinh nội dung tạo bố cục hoàn chỉnh; việc gửi lại khi lỗi dịch vụ cần có giới hạn và lưu từng lượt, không dùng vòng gọi model không kiểm soát."),
    ],
    [
        ("Heading3", "2.7.3. Kết quả thử nghiệm và giới hạn"),
        ("Normal", "Thử nghiệm ngày 12/10/2026 sử dụng ảnh kế hoạch thực tế đã lưu trong dự án. Bộ chuẩn bị dữ liệu đọc lại kết quả PP-StructureV3 trên cùng ảnh, không chạy nhận diện Paddle mới. Cả ba cấu hình native, OCRv5 và OCRv6 có bảng HTML 4 cột, 8 hàng tổng; 7 hàng sau hàng tiêu đề đều trống. Dữ liệu gửi Gemini gồm 32 dòng OCR nguyên trạng, HTML bảng và tọa độ ô của cấu hình native."),
        ("Normal", "Hai request độc lập dùng Gemini 3.8 Flash, prompt tiếng Việt, temperature 0 và mức thinking medium. Request thứ nhất trả HTTP 200, sinh nội dung trong 17,94 giây; tổng gồm xuất Word và render local là 19,50 giây. API ghi nhận 11.594 token. Word tạo được có một trang, bảng thực 4 cột và 7 hàng thân trống, không nhúng ảnh hoặc in ra ký tự escape xuống dòng. Request thứ hai nhận HTTP 503 sau 6,31 giây, không tạo Word."),
        ("Normal", "Số hàng và cột của bản thành công đã đúng với ảnh. Tuy nhiên, ghi chú và hướng dẫn bị gộp đoạn, thiếu dấu đóng ngoặc trong ghi chú; cỡ chữ, giãn cách và in đậm đề tài còn khác nguồn. Dấu đỏ không được tái dựng trong thân tài liệu. Bản tự động được giữ nguyên khi đánh giá, không sửa tay để làm tăng kết quả kiểm thử."),
        ("Normal", "Lượt này đồng thời thay ngôn ngữ prompt và bổ sung cấu trúc bảng Paddle, nên chưa xác định riêng mức đóng góp của từng thay đổi. Kết quả trên một mẫu chưa chứng minh độ ổn định với nhiều loại công văn; chưa đo CER/WER hoặc tỷ lệ giữ đúng bố cục trên bộ dữ liệu lớn. Các số thời gian trên không phải cam kết hiệu năng triển khai."),
        ("Heading3", "2.7.4. Phạm vi triển khai và hướng hoàn thiện"),
        ("Normal", "Phương án đã được lựa chọn và kiểm thử bằng script; chưa tích hợp pipeline bố cục mới vào endpoint Word của ứng dụng. Công việc tiếp theo là tích hợp OCRService, theo dõi trạng thái tác vụ, lưu bản gốc/bản Word và phiên bản chỉnh sửa, xử lý lỗi API, đồng thời cho văn thư xem ảnh cạnh kết quả. Trình soạn thảo Word trong ứng dụng là hạng mục cần triển khai, không phải chức năng đã nghiệm thu."),
        ("Normal", "Bằng chứng: testcase TC-OCR-GEMINI-PADDLE-VI-062 trong tiendo/7_test_cases.md; dữ liệu Paddle, response, layout, metadata và Word tại tai_nguyen/ket_qua_ocr/gemini38_vi_paddle_table. Hướng dẫn kỹ thuật tham khảo: tài liệu PP-StructureV3 của PaddleOCR và trang model Gemini 3.8 Flash của Google AI for Developers (truy cập ngày 12/10/2026)."),
        ("Reference", "PaddleOCR: https://www.paddleocr.ai/main/en/version3.x/pipeline_usage/PP-StructureV3.html"),
        ("Reference", "Google: https://ai.google.dev/gemini-api/docs/models/gemini-3.8-flash"),
    ],
]


def section_xml():
    doc = Document()
    for page_index, page in enumerate(PAGES):
        for index, (style, text) in enumerate(page):
            paragraph = doc.add_paragraph(style=style if style != "Reference" else "Normal")
            fmt = paragraph.paragraph_format
            fmt.page_break_before = index == 0 or text.startswith("2.7.4.")
            fmt.space_before = Pt(6 if style.startswith("Heading") else 0)
            fmt.space_after = Pt(6 if style.startswith("Heading") else 3)
            fmt.line_spacing = 1.1
            fmt.keep_with_next = style.startswith("Heading")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if style.startswith("Heading") or style == "Reference" else WD_ALIGN_PARAGRAPH.JUSTIFY
            run = paragraph.add_run(text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(14 if style != "Reference" else 10)
            run.font.color.rgb = None
            run.bold = style.startswith("Heading")
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "000000")
            run._element.get_or_add_rPr().append(color)
    return b"".join(etree.tostring(p) for p in doc.element.body if p.tag != qn("w:sectPr"))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if args.output.exists():
        raise ValueError("Output exists; do not overwrite a reviewed report")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    inserted = section_xml()
    with ZipFile(args.source) as source, ZipFile(args.output, "w") as target:
        original = source.read("word/document.xml")
        marker = original.rfind(b"<w:sectPr")
        if marker < 0 or original.rfind(b"</w:p>") > marker:
            raise ValueError("Expected final body section properties")
        updated = original[:marker] + inserted + original[marker:]
        etree.fromstring(updated)
        for entry in source.infolist():
            target.writestr(entry, updated if entry.filename == "word/document.xml" else source.read(entry.filename))
    with ZipFile(args.output) as updated_zip, ZipFile(args.source) as original_zip:
        unchanged = all(updated_zip.read(name) == original_zip.read(name)
                        for name in original_zip.namelist() if name != "word/document.xml")
    metadata = {"source_sha256": hashlib.sha256(args.source.read_bytes()).hexdigest(),
                "output_sha256": hashlib.sha256(args.output.read_bytes()).hexdigest(),
                "existing_document_xml_preserved": updated.replace(inserted, b"", 1) == original,
                "other_package_parts_preserved": unchanged, "added_section": "2.7", "model_calls": 0}
    args.output.with_suffix(".verification.json").write_text(json.dumps(metadata, indent=2), encoding="utf-8")
    print(json.dumps(metadata))


if __name__ == "__main__":
    main()
